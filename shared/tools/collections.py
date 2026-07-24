"""
Shared collection-introspection tools for BV-BRC Solr.

These tools use data parsed from ``data_types.xlsx`` rather than the MCP
server's prompt-file-based implementations.  The spreadsheet provides
richer metadata (field types, definitions, examples).

Any agent can import ``list_collections`` / ``get_collection_fields`` to
discover what data is available.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Dict, List, Optional

# ---------------------------------------------------------------------------
# Load collection metadata from the spreadsheet (cached at module level)
# ---------------------------------------------------------------------------

_COLLECTION_DATA: Optional[Dict[str, Dict[str, Any]]] = None
_COLLECTION_INDEX: Optional[List[Dict[str, str]]] = None


def _load_collection_data() -> tuple[list[dict[str, str]], dict[str, dict[str, Any]]]:
    """Parse ``data_types.xlsx`` and cache the results.

    Returns ``(index, collections)`` where *index* is a list of
    ``{name, purpose}`` dicts and *collections* maps collection name to
    ``{primary_key, fields: [{name, type, definition}]}``.
    """
    global _COLLECTION_DATA, _COLLECTION_INDEX

    if _COLLECTION_DATA is not None and _COLLECTION_INDEX is not None:
        return _COLLECTION_INDEX, _COLLECTION_DATA

    try:
        import openpyxl
    except ImportError:
        _COLLECTION_INDEX = []
        _COLLECTION_DATA = {}
        return _COLLECTION_INDEX, _COLLECTION_DATA

    xlsx_path = _find_spreadsheet()
    if xlsx_path is None:
        _COLLECTION_INDEX = []
        _COLLECTION_DATA = {}
        return _COLLECTION_INDEX, _COLLECTION_DATA

    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True)

    # Parse index sheet
    index: list[dict[str, str]] = []
    ws_index = wb["data_types"]
    seen_names: set[str] = set()
    for row in ws_index.iter_rows(min_row=2, values_only=True):
        name = row[0]
        purpose = row[1]
        if name and name not in seen_names:
            seen_names.add(name)
            index.append({"name": str(name).strip(), "purpose": str(purpose).strip()})

    # Excluded system fields
    excluded = {
        "version",
        "owner",
        "public",
        "user_read",
        "user_write",
        "date_inserted",
        "date_modified",
        "_version_",
    }

    # Parse each collection sheet
    collections: Dict[str, Dict[str, Any]] = {}
    for sheet_name in wb.sheetnames:
        if sheet_name == "data_types":
            continue

        ws = wb[sheet_name]
        metadata: Dict[str, str] = {}
        fields: List[Dict[str, str]] = []
        in_fields = False

        for row in ws.iter_rows(values_only=True):
            vals = list(row)

            if vals[0] in ("Data type", "Primary key", "Purpose"):
                metadata[str(vals[0])] = str(vals[1]).strip() if vals[1] else ""
                continue

            if vals[0] == "Attribute name":
                in_fields = True
                continue

            if in_fields and vals[0] is not None:
                field_name = str(vals[0]).strip().rstrip("*")
                if field_name.lower() in excluded:
                    continue

                field_type = str(vals[1]).strip() if vals[1] else "string"
                definition = str(vals[2]).strip() if vals[2] else ""

                fields.append(
                    {
                        "name": field_name,
                        "type": field_type,
                        "definition": definition,
                    }
                )

        collections[sheet_name] = {
            "primary_key": metadata.get("Primary key", ""),
            "purpose": metadata.get("Purpose", ""),
            "fields": fields,
        }

    wb.close()

    _COLLECTION_INDEX = index
    _COLLECTION_DATA = collections
    return _COLLECTION_INDEX, _COLLECTION_DATA


def _find_spreadsheet() -> Optional[Path]:
    """Locate ``data_types.xlsx``."""
    candidates = [
        # Relative to the agents/ directory
        Path(__file__).resolve().parent.parent.parent / "agents" / "data_types.xlsx",
        # Legacy location
        Path(__file__).resolve().parent.parent.parent / "data_types.xlsx",
        Path.cwd() / "data_types.xlsx",
    ]
    for p in candidates:
        if p.exists():
            return p
    return None


# ---------------------------------------------------------------------------
# Tool functions
# ---------------------------------------------------------------------------


async def list_collections(**kwargs: Any) -> Dict[str, Any]:
    """List all available BV-BRC Solr collections with descriptions.

    Returns:
        Dict with ``collections`` (list of ``{name, purpose}``) and ``count``.
    """
    index, _ = _load_collection_data()

    if not index:
        return {
            "collections": [],
            "count": 0,
            "note": "Collection metadata unavailable (data_types.xlsx not found).",
        }

    return {
        "collections": index,
        "count": len(index),
    }


async def get_collection_fields(collection: str, **kwargs: Any) -> Dict[str, Any]:
    """Get the queryable fields for a specific BV-BRC Solr collection.

    Args:
        collection: The collection name to inspect.

    Returns:
        Dict with ``fields`` (list of ``{name, type, definition}``),
        or an error if the collection is not found.
    """
    _, collections = _load_collection_data()

    cdata = collections.get(collection)
    if cdata is None:
        return {
            "error": f"Collection '{collection}' not found.",
            "available_collections": sorted(collections.keys()),
        }

    return {
        "collection": collection,
        "primary_key": cdata["primary_key"],
        "field_count": len(cdata["fields"]),
        "fields": cdata["fields"],
    }
